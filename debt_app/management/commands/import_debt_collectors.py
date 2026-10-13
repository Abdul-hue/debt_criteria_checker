"""
Load debt collectors (BPO) into the DebtCollector list used by the BPO-DEBT rule.

    python manage.py import_debt_collectors collectors.csv [--dry-run]

Accepts .csv, .txt or .xlsx (first sheet). Column 1 is the collector name;
optional column 2 holds other trading names separated by ";". A header row
whose first cell is "name" is skipped.

Additive only: new names are created, existing names get any new aliases and
are re-activated. Nothing is deleted or deactivated — do that in the admin.
"""
import csv
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from debt_app.models import DebtCollector


def _read_rows(path: Path) -> list:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        return [[("" if v is None else str(v)) for v in row] for row in sheet.iter_rows(values_only=True)]
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.reader(f))


class Command(BaseCommand):
    help = "Import debt collectors (BPO) from a .csv/.txt/.xlsx file into the DebtCollector list."

    def add_arguments(self, parser):
        parser.add_argument("file")
        parser.add_argument("--dry-run", action="store_true", help="Show what would change without saving.")

    def handle(self, *args, **options):
        path = Path(options["file"])
        if not path.exists():
            raise CommandError(f"File not found: {path}")

        created, updated, unchanged = [], [], []
        with transaction.atomic():
            for row in _read_rows(path):
                name = (row[0] if row else "").strip()
                if not name or name.lower() == "name":
                    continue
                aliases = [a.strip() for a in (row[1] if len(row) > 1 else "").split(";") if a.strip()]

                obj = DebtCollector.objects.filter(name__iexact=name).first()
                if obj is None:
                    DebtCollector.objects.create(name=name, aliases=aliases)
                    created.append(name)
                    continue
                merged = list(obj.aliases or [])
                merged += [a for a in aliases if a.lower() not in {m.lower() for m in merged}]
                if merged != (obj.aliases or []) or not obj.is_active:
                    obj.aliases, obj.is_active = merged, True
                    obj.save()
                    updated.append(name)
                else:
                    unchanged.append(name)

            if options["dry_run"]:
                transaction.set_rollback(True)

        prefix = "[DRY RUN] " if options["dry_run"] else ""
        for name in created:
            self.stdout.write(f"  {prefix}[CREATED] {name}")
        for name in updated:
            self.stdout.write(f"  {prefix}[UPDATED] {name}")
        self.stdout.write(self.style.SUCCESS(
            f"{prefix}{len(created)} created, {len(updated)} updated, {len(unchanged)} unchanged."
        ))
