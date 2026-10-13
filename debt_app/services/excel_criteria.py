"""
Turn the criteria workbook ("TIP CRITERIA & VOTING HISTORY.xlsx") into the
markdown files under ``Excel Criteria/`` that the seed commands parse, and
compare two versions of the workbook.

The markdown files are what the app actually reads (no seeder opens the
workbook), so this module makes that conversion repeatable and reviewable in
git instead of a one-off hand job.

Two deliberate behaviours:

* **Known data fixes** (``GENERAL_ROW_FIXES``, ``COUNCIL_ROW_DROPS``,
  ``COUNCIL_NAME_FIXES``) correct mistakes in the workbook that the owner has
  acknowledged. Each fix only fires when the cell still holds the expected
  wrong value, and every firing is reported, so a corrected workbook makes the
  fix a no-op rather than a surprise.
* **Kept rows.** The markdown files have accumulated hand additions that are
  not in the workbook (creditors restored after a destructive seed run,
  representative mappings added by migrations). A regeneration keeps anything
  the existing file lists that the workbook does not, appends it under a
  marker, and reports it, so nothing silently disappears. ``drop_missing``
  turns that off.

The prose criteria sheets (TIG / Watch / TIX) are hand-maintained markdown and
are not generated here.
"""

import datetime
import re
from pathlib import Path

# --------------------------------------------------------------------------
# Workbook layout
# --------------------------------------------------------------------------

SHEET_GENERAL = "GENERAL CREDITOR"
SHEET_COUNCILS = "Councils"
SHEET_WHICH_REP = "Which Representative"
SHEET_DIVIDENDS = "Dividends "  # the trailing space is in the workbook
SHEET_COUNTY = "County Councils"
SHEET_BANKING = "Banking Groups"
SHEET_TIG = "TIG Criteria"
SHEET_WATCH = "Watch Criteria"
SHEET_TIX = "TIX Criteria"
SHEET_EVOLVE = "Evolve Criteria"

FILE_GENERAL = "General_Creditors.md"
FILE_COUNCILS = "Councils.md"
FILE_WHICH_REP = "Which_Representative_Criteria.md"
FILE_DIVIDENDS = "Dividends_Criteria.md"
FILE_COUNTY = "County_Councils_Criteria.md"
FILE_BANKING = "Banking_Groups_Criteria.md"

GENERAL_COLS = 7   # Creditor, Accept/Reject, Notes, Updated, Phone, Email, Contact
COUNCIL_COLS = 7   # Council, Accept/Reject, Notes, Updated, Changed-from-Rej date, Contact Name, Contact Number
COUNTY_COLS = 4    # Council, Accept/Reject, Notes, Districts
DIVIDEND_COLS = 3  # Creditor, Div Required, Notes

GENERAL_HEADER = "| Creditor | Accept / Reject | Notes / Criteria | updated criteria | Phone Number | Email | Contact |"
GENERAL_SEPARATOR = "| --- | --- | --- | --- | --- | --- | --- |"
# seed_all_councils looks for these two lines verbatim.
COUNCIL_HEADER = "| # | Council | Status | Notes | Updated Criteria | Criteria Changed From Rej Date | Contact Name | Contact Number |"
COUNCIL_SEPARATOR = "|---|---------|--------|-------|------------------|-------------------------------|--------------|----------------|"
DIVIDEND_HEADER = "| Creditor | Div Required | Notes |"
DIVIDEND_SEPARATOR = "| --- | --- | --- |"

# Order the representative sections are written in (matches the old file).
REP_SECTION_ORDER = ["(WATCH) WPM", "TIX", "EVOLVE", "EVERYDAY LOANS"]
REP_SECTION_EVOLVE = "EVOLVE"

KEPT_MARKER = "<!-- The rows below are not in the workbook. They were kept from the previous version of this file. Delete a row here to retire it. -->"
KEPT_SUBHEADING = "### Kept from the previous file (not in the workbook)"

# --------------------------------------------------------------------------
# Known data fixes (acknowledged by the sheet owner, 2026-10-12)
# --------------------------------------------------------------------------

GENERAL_ROW_FIXES = [
    {
        "name": "startline motor finance",
        "col": 1,  # Accept / Reject
        "expect": "can get a vote if chased",
        "value": "DO NOT VOTE",
        "move_old_to_notes": True,
        "why": "'can get a vote if chased' is not a vote status; kept as DO NOT VOTE with the wording moved to the notes",
    },
    {
        "name": "partners credit union",
        "col": 3,  # updated criteria
        "expect": "03/12/2026",
        "value": "12/03/2026",
        "move_old_to_notes": False,
        "why": "day and month were swapped; the note says confirmed 12/03/2026",
    },
]

# Stances the sheet owner gave by email on 2026-10-13 for rows whose
# 'Accept / Reject' cell was blank, a remark, or "unsure". (name, cell as the
# workbook has it, stance). Rows she said to leave until a case shows how they
# vote (Capital on Tap, Leap Utilites, Clonard CU, Hull University, RIA, Kroo,
# Peabody, Western Circle, Reevo, Claims Advisory Group, Glasgow CU, Malden
# Housing) are deliberately absent: they keep their current database status.
OWNER_ANSWERS_2026_10_13 = [
    ("reward rate", "trading style of amigo", "DO NOT VOTE"),
    ("24/7 moneybox / active securities", "", "DO NOT VOTE"),
    ("a shade greener boilers", "", "REJECT"),
    ("1st stop personal loans / oplo", "", "REJECT"),
    ("billings finance", "will consider", "ACCEPT"),
    ("365 business finance ltd", "", "REJECT"),
    ("3 mobile", "", "DO NOT VOTE"),
    ("3j finance", "", "DO NOT VOTE"),
    ("creditspring", "", "DO NOT VOTE"),
    ("sll capital", "croydon", "ACCEPT"),
]
GENERAL_ROW_FIXES += [
    {
        "name": name,
        "col": 1,  # Accept / Reject
        "expect": expect,
        "value": value,
        "move_old_to_notes": bool(expect),
        "why": f"stance confirmed by the sheet owner on 2026-10-13: {value}",
    }
    for name, expect, value in OWNER_ANSWERS_2026_10_13
]

COUNCIL_ROW_DROPS = [
    {
        "name": "ashford borough council",
        "col": 1,  # Accept / Reject
        "expect": "do not vote",
        "why": "this second Ashford row is Aylesbury Vale's record with the name overwritten; Ashford keeps its REJECT row",
    },
    {
        "name": "wycombe district council",
        "col": 1,  # Accept / Reject
        "expect": "will consider",
        "why": "duplicate Wycombe row with no notes; the sheet owner confirmed on 2026-10-13 it can be deleted and the Accept row stands",
    },
]

COUNCIL_NAME_FIXES = {
    "copeland borough council forms part of cumberland council as of jan 2026": (
        "Copeland Borough Council",
        "Forms part of Cumberland Council as of Jan 2026.",
    ),
}

# Name cells that had a remark typed into them. Key = the cell as the workbook
# has it (lower-cased, whitespace collapsed); value = (real name, remark that
# goes to the front of the notes). Keeps the name matching the database row.
GENERAL_NAME_FIXES = {
    "match the cash t/a guarantor my loan (match the cash trading name) also trader as 'sharemyloan' share my loan'": (
        "Match the Cash t/a Guarantor My Loan (Match the Cash trading name)",
        "Also trades as 'Sharemyloan' / Share my loan.",
    ),
    "hillingdon credit union - see london community credit union": (
        "Hillingdon Credit Union",
        "See London Community Credit Union.",
    ),
}

# Rows the workbook genuinely renamed. Key = name as the previous file (and the
# database row) has it; value = the name the workbook uses now. The exporter
# does not keep the old row, seed_creditor_criteria renames the existing
# database row (same primary key) instead of creating a second one, and
# build_criteria_change_from_excel proposes the old name as a trading name so
# cases still naming the creditor the old way keep resolving.
RENAMED_ROWS = {
    "Clever Money (Blackpool, Fylde & Wyre Credit Union Ltd)": "Clevr Money (Blackpool, Fylde & Wyre Credit Union Ltd)",
    "Knowsley Credit Union": "Knowsley Mutual Credit Union",
}
_RENAMED_LOWER = {k.lower(): (k, v) for k, v in RENAMED_ROWS.items()}

# Junk rows the previous files carry that must not be kept on regeneration.
KEPT_ROW_IGNORE = {"rth"}

# County rows whose name cell is blank: recognise them by their first district.
BLANK_COUNTY_BY_FIRST_DISTRICT = {
    "sevenoaks": "Kent County Council",
    "norwich": "Norfolk County Council",
}

# --------------------------------------------------------------------------
# Cell helpers
# --------------------------------------------------------------------------

_WS = re.compile(r"[\s\xa0  ]+")
_LINE_BREAKS = re.compile(r"[\r\n  ]")
_DISTRICT_SPLIT = re.compile(r"\s{3,}|\r?\n")


def clean_text(value):
    """Cell value -> single-line text.

    None -> '', dates -> DD/MM/YYYY, whole floats -> ints, all whitespace runs
    (including non-breaking spaces and newlines) collapsed to one space, and
    pipes replaced so a cell can never break the markdown table.
    """
    if value is None:
        return ""
    if isinstance(value, datetime.date):  # datetime.datetime is a subclass
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("|", "/")
    return _WS.sub(" ", text).strip()


def clean_name(value):
    """Name cell -> text, keeping internal spacing exactly as the workbook has it.

    Creditor and council names are matched verbatim against database rows that
    were seeded from earlier exports, and some of those names carry doubled or
    non-breaking spaces ("Omno  Capital", "Isle of Anglesey"). Collapsing
    them would turn a reseed into a duplicate row, so only line breaks, pipes
    and leading/trailing whitespace are touched.
    """
    if value is None:
        return ""
    text = str(value).replace("|", "/")
    # One space per line break (not per run): "North Yorkshire.\n\n\nNorth" was
    # seeded as "North Yorkshire.   North" and must still match that row.
    text = _LINE_BREAKS.sub(" ", text.replace("\r\n", "\n"))
    return text.strip()


def _lower(text):
    return clean_text(text).lower()


def load_workbook(path):
    from openpyxl import load_workbook as _load

    return _load(str(path), data_only=True, read_only=False)


def read_rows(ws, ncols, first_row=2):
    """[(row_number, [raw cells...])] for every row with at least one value."""
    rows = []
    for r in range(first_row, ws.max_row + 1):
        raw = [ws.cell(r, c).value for c in range(1, ncols + 1)]
        if any(clean_text(v) for v in raw):
            rows.append((r, raw))
    return rows


def split_districts(raw):
    """The Districts cell holds one line per district separated by runs of
    spaces or newlines. Return the cleaned list."""
    if raw is None:
        return []
    items = []
    for part in _DISTRICT_SPLIT.split(str(raw).replace("\xa0", " ")):
        part = _WS.sub(" ", part).strip().rstrip("-").strip()
        if part:
            items.append(part)
    return items


# --------------------------------------------------------------------------
# Existing-file readers (for kept rows)
# --------------------------------------------------------------------------

def _is_separator(cells):
    return all(re.fullmatch(r":?-+:?", c or "-") for c in cells)


def existing_table_rows(path, ncols, numbered=False):
    """Data rows of a markdown pipe table, padded/trimmed to ``ncols``."""
    path = Path(path)
    if not path.exists():
        return []
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if numbered:
            cells = cells[1:]
        if not cells or _is_separator(cells):
            continue
        if cells[0] in ("Creditor", "Council", "#"):
            continue
        cells = (cells + [""] * ncols)[:ncols]
        rows.append(cells)
    return rows


def existing_rep_sections(path):
    """{heading (as written): [items]} from Which_Representative_Criteria.md.
    '### ' sub-headings are transparent: their items belong to the enclosing
    '## ' section, which is also how seed_creditor_criteria reads the file."""
    path = Path(path)
    sections = {}
    if not path.exists():
        return sections
    current = None
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):
            current = line[3:].strip()
            sections.setdefault(current, [])
        elif line.startswith("- ") and current is not None:
            item = line[2:].strip()
            if item:
                sections[current].append(item)
    return sections


# --------------------------------------------------------------------------
# Report
# --------------------------------------------------------------------------

class ExportReport:
    """Everything a reviewer needs to know about one export run."""

    def __init__(self):
        self.files = {}        # filename -> {"rows": int, "kept": [..], "fixes": [..], "dropped": [..], "notes": [..]}
        self.warnings = []

    def file(self, name):
        return self.files.setdefault(
            name, {"rows": 0, "kept": [], "fixes": [], "dropped": [], "renamed": [], "notes": []}
        )

    def warn(self, message):
        self.warnings.append(message)

    def as_text(self):
        out = []
        for name, info in self.files.items():
            out.append(f"{name}: {info['rows']} row(s) written")
            for f in info["fixes"]:
                out.append(f"    FIX     {f}")
            for d in info["dropped"]:
                out.append(f"    DROPPED {d}")
            for rn in info["renamed"]:
                out.append(f"    RENAMED {rn}")
            for k in info["kept"]:
                out.append(f"    KEPT    {k}")
            for n in info["notes"]:
                out.append(f"    NOTE    {n}")
        for w in self.warnings:
            out.append(f"WARNING {w}")
        return "\n".join(out)


# --------------------------------------------------------------------------
# Exporter
# --------------------------------------------------------------------------

def _source_line(xlsx_name, sheet, today):
    return (
        f"> Source: `{xlsx_name}` — sheet `{sheet}` — exported {today:%Y-%m-%d} by "
        f"`manage.py export_excel_criteria`. Regenerate with that command rather than "
        f"editing rows by hand; rows added by hand are kept on regeneration and listed at the end."
    )


def _table_line(cells):
    return "| " + " | ".join(cells) + " |"


def _apply_name_fix(cells, fixes, info, row_number):
    """If the name cell is a known 'name + remark' cell, restore the real name
    and move the remark to the front of the notes column."""
    fix = fixes.get(_lower(cells[0]))
    if fix:
        old = cells[0]
        cells[0], remark = fix
        cells[2] = f"{remark} {cells[2]}".strip()
        info["fixes"].append(f"row {row_number}: name {old!r} -> {cells[0]!r}; wording moved to notes")


def _kept_rows(existing_rows, exported_names, info):
    """Rows of the previous file the workbook no longer mentions, minus
    renamed rows (reported) and junk rows."""
    kept = []
    for row in existing_rows:
        name = row[0]
        if not name or name.lower() in exported_names:
            continue
        key = _lower(name)
        if key in KEPT_ROW_IGNORE:
            continue
        if key in GENERAL_NAME_FIXES or key in COUNCIL_NAME_FIXES:
            # An uncorrected "name + remark" cell left by an earlier export;
            # the corrected row is already in the output.
            info["notes"].append(f"{name!r} is a stale annotated name from a previous export; not kept")
            continue
        if key in _RENAMED_LOWER:
            info["renamed"].append(f"{name!r} -> {_RENAMED_LOWER[key][1]!r} (old row not kept; the database sync renames it)")
            continue
        kept.append(row)
    return kept


def export_workbook(xlsx_path, out_dir, *, drop_missing=False, keep_evolve=False,
                    today=None, write=True):
    """Write the markdown files for ``xlsx_path`` into ``out_dir``.

    Returns (ExportReport, {filename: content}). With ``write=False`` nothing
    touches disk, which is what ``--dry-run`` uses.
    """
    xlsx_path = Path(xlsx_path)
    out_dir = Path(out_dir)
    today = today or datetime.date.today()
    wb = load_workbook(xlsx_path)
    report = ExportReport()
    outputs = {}

    def src(sheet):
        return _source_line(xlsx_path.name, sheet, today)

    # ---- GENERAL CREDITOR ------------------------------------------------
    if SHEET_GENERAL in wb.sheetnames:
        info = report.file(FILE_GENERAL)
        lines = ["# GENERAL CREDITOR", "", src(SHEET_GENERAL), "", GENERAL_HEADER, GENERAL_SEPARATOR]
        names = set()
        for r, raw in read_rows(wb[SHEET_GENERAL], GENERAL_COLS):
            cells = [clean_name(raw[0])] + [clean_text(v) for v in raw[1:]]
            _apply_name_fix(cells, GENERAL_NAME_FIXES, info, r)
            for fix in GENERAL_ROW_FIXES:
                if _lower(cells[0]) == fix["name"] and _lower(cells[fix["col"]]) == _lower(fix["expect"]):
                    old = cells[fix["col"]]
                    cells[fix["col"]] = fix["value"]
                    if fix["move_old_to_notes"]:
                        cells[2] = f"Sheet says: '{old}'. " + cells[2] if cells[2] else f"Sheet says: '{old}'."
                    info["fixes"].append(f"row {r} {cells[0]!r}: {old!r} -> {fix['value']!r} ({fix['why']})")
            if cells[0]:
                names.add(cells[0].lower())
            lines.append(_table_line(cells))
            info["rows"] += 1
        if not drop_missing:
            kept = _kept_rows(existing_table_rows(out_dir / FILE_GENERAL, GENERAL_COLS), names, info)
            if kept:
                lines += ["", KEPT_MARKER]
                for row in kept:
                    lines.append(_table_line(row))
                    info["kept"].append(f"{row[0]!r} ({row[1] or 'no status'})")
                info["rows"] += len(kept)
        outputs[FILE_GENERAL] = "\n".join(lines) + "\n"
    else:
        report.warn(f"sheet {SHEET_GENERAL!r} not found; {FILE_GENERAL} not written")

    # ---- Councils ----------------------------------------------------------
    if SHEET_COUNCILS in wb.sheetnames:
        info = report.file(FILE_COUNCILS)
        body = []
        names = set()
        for r, raw in read_rows(wb[SHEET_COUNCILS], COUNCIL_COLS):
            cells = [clean_name(raw[0])] + [clean_text(v) for v in raw[1:]]
            dropped = False
            for drop in COUNCIL_ROW_DROPS:
                if _lower(cells[0]) == drop["name"] and _lower(cells[drop["col"]]) == drop["expect"]:
                    info["dropped"].append(f"row {r} {cells[0]!r} ({cells[1]!r}): {drop['why']}")
                    dropped = True
            if dropped:
                continue
            _apply_name_fix(cells, COUNCIL_NAME_FIXES, info, r)
            if cells[0]:
                names.add(cells[0].lower())
            body.append(cells)
        kept_rows = []
        if not drop_missing:
            kept_rows = _kept_rows(
                existing_table_rows(out_dir / FILE_COUNCILS, COUNCIL_COLS, numbered=True), names, info
            )
        lines = [
            "# Councils Sheet - TIP CRITERIA and VOTING HISTORY", "",
            src(SHEET_COUNCILS),
            f"> Total entries: {len(body) + len(kept_rows)}", "",
            COUNCIL_HEADER, COUNCIL_SEPARATOR,
        ]
        n = 0
        for cells in body:
            n += 1
            lines.append(_table_line([str(n)] + cells))
        info["rows"] = n
        if kept_rows:
            lines += ["", KEPT_MARKER]
            for row in kept_rows:
                n += 1
                lines.append(_table_line([str(n)] + row))
                info["kept"].append(f"{row[0]!r} ({row[1] or 'no status'})")
            info["rows"] = n
        outputs[FILE_COUNCILS] = "\n".join(lines) + "\n"
    else:
        report.warn(f"sheet {SHEET_COUNCILS!r} not found; {FILE_COUNCILS} not written")

    # ---- Which Representative ---------------------------------------------
    if SHEET_WHICH_REP in wb.sheetnames:
        info = report.file(FILE_WHICH_REP)
        ws = wb[SHEET_WHICH_REP]
        sheet_sections = {}  # heading -> [items]
        for c in range(1, ws.max_column + 1):
            heading = clean_text(ws.cell(1, c).value)
            if not heading:
                continue
            items = []
            for r in range(2, ws.max_row + 1):
                item = clean_name(ws.cell(r, c).value)
                if item:
                    items.append(item)
            # The sheet repeats some names; the file has always been a sorted,
            # de-duplicated list, which also keeps its git diff readable.
            sheet_sections[heading] = sorted(dict.fromkeys(items), key=str.lower)

        existing = existing_rep_sections(out_dir / FILE_WHICH_REP)
        existing_lower = {h.lower(): items for h, items in existing.items()}

        ordered = [h for h in REP_SECTION_ORDER if h in sheet_sections]
        ordered += [h for h in sheet_sections if h not in ordered]

        lines = ["# Which Representative", "", src(SHEET_WHICH_REP)]
        for heading in ordered:
            items = sheet_sections[heading]
            if heading.upper() == REP_SECTION_EVOLVE and not keep_evolve:
                prev = existing_lower.get(heading.lower(), [])
                info["dropped"].append(
                    f"section {heading!r}: {len(items)} workbook entries and "
                    f"{len([p for p in prev if p not in items])} previously kept entries "
                    f"(Evolve no longer exists; confirmed by the sheet owner 2026-10-12)"
                )
                continue
            lines += ["", f"## {heading}"]
            for item in items:
                lines.append(f"- {item}")
            info["rows"] += len(items)
            if not drop_missing:
                kept = [p for p in existing_lower.get(heading.lower(), []) if p not in items]
                if kept:
                    lines += ["", KEPT_SUBHEADING]
                    for item in kept:
                        lines.append(f"- {item}")
                        info["kept"].append(f"{heading}: {item!r}")
                    info["rows"] += len(kept)
        outputs[FILE_WHICH_REP] = "\n".join(lines) + "\n"
    else:
        report.warn(f"sheet {SHEET_WHICH_REP!r} not found; {FILE_WHICH_REP} not written")

    # ---- Dividends ---------------------------------------------------------
    if SHEET_DIVIDENDS in wb.sheetnames:
        info = report.file(FILE_DIVIDENDS)
        lines = ["# Dividends", "", src(SHEET_DIVIDENDS), "", DIVIDEND_HEADER, DIVIDEND_SEPARATOR]
        names = set()
        for r, raw in read_rows(wb[SHEET_DIVIDENDS], DIVIDEND_COLS):
            cells = [clean_name(raw[0])] + [clean_text(v) for v in raw[1:]]
            if not cells[0]:
                continue
            names.add(cells[0].lower())
            lines.append(_table_line(cells))
            info["rows"] += 1
        if not drop_missing:
            kept = _kept_rows(existing_table_rows(out_dir / FILE_DIVIDENDS, DIVIDEND_COLS), names, info)
            if kept:
                lines += ["", KEPT_MARKER]
                for row in kept:
                    lines.append(_table_line(row))
                    info["kept"].append(f"{row[0]!r} ({row[1] or 'no dividend'})")
                info["rows"] += len(kept)
        outputs[FILE_DIVIDENDS] = "\n".join(lines) + "\n"
    else:
        report.warn(f"sheet {SHEET_DIVIDENDS!r} not found; {FILE_DIVIDENDS} not written")

    # ---- County Councils ---------------------------------------------------
    if SHEET_COUNTY in wb.sheetnames:
        info = report.file(FILE_COUNTY)
        lines = ["# County Councils", "", src(SHEET_COUNTY)]
        for r, raw in read_rows(wb[SHEET_COUNTY], COUNTY_COLS):
            name = clean_name(raw[0])
            accept_reject, notes = (clean_text(v) for v in raw[1:3])
            districts = split_districts(raw[3])
            if not name:
                first = districts[0].split()[0].lower() if districts else ""
                name = BLANK_COUNTY_BY_FIRST_DISTRICT.get(first)
                if name:
                    info["fixes"].append(f"row {r}: blank county name resolved to {name!r} from its districts")
                else:
                    name = f"Unnamed county (workbook row {r})"
                    report.warn(f"{SHEET_COUNTY} row {r} has no name and its districts are not recognised")
            lines += ["", f"## {name}", ""]
            if accept_reject:
                lines.append(f"**Accept / Reject:** {accept_reject}  ")
            lines.append(f"**Notes:** {notes}  ")
            lines += ["", "**Districts:**"]
            for d in districts:
                lines.append(f"- {d}")
            info["rows"] += 1
        outputs[FILE_COUNTY] = "\n".join(lines) + "\n"
    else:
        report.warn(f"sheet {SHEET_COUNTY!r} not found; {FILE_COUNTY} not written")

    # ---- Banking Groups ----------------------------------------------------
    if SHEET_BANKING in wb.sheetnames:
        info = report.file(FILE_BANKING)
        ws = wb[SHEET_BANKING]
        lines = ["# Banking Groups", "", src(SHEET_BANKING)]
        for c in range(1, ws.max_column + 1):
            heading = clean_text(ws.cell(1, c).value)
            if not heading:
                continue
            lines += ["", f"## {heading}", ""]
            for r in range(2, ws.max_row + 1):
                item = clean_name(ws.cell(r, c).value)
                if item:
                    lines.append(f"- {item}")
                    info["rows"] += 1
        outputs[FILE_BANKING] = "\n".join(lines) + "\n"
    else:
        report.warn(f"sheet {SHEET_BANKING!r} not found; {FILE_BANKING} not written")

    if SHEET_EVOLVE in wb.sheetnames:
        report.warn(f"workbook still has an {SHEET_EVOLVE!r} sheet; Evolve was confirmed retired on 2026-10-12")

    if write:
        out_dir.mkdir(parents=True, exist_ok=True)
        for name, content in outputs.items():
            (out_dir / name).write_text(content, encoding="utf-8")

    return report, outputs


# --------------------------------------------------------------------------
# Workbook comparison
# --------------------------------------------------------------------------

def _keyed_rows(ws, ncols):
    """{name.lower(): [cells...]} for a name-keyed sheet; first occurrence wins,
    extra occurrences are counted so renames and duplicates show up."""
    keyed, counts = {}, {}
    for _r, raw in read_rows(ws, ncols):
        cells = [clean_text(v) for v in raw]
        key = cells[0].lower()
        if not key:
            continue
        counts[key] = counts.get(key, 0) + 1
        keyed.setdefault(key, cells)
    return keyed, counts


def _short(text, n=160):
    text = str(text)
    return text if len(text) <= n else text[: n - 3] + "..."


def _text_cells(ws):
    return [clean_text(c.value) for row in ws.iter_rows() for c in row if clean_text(c.value)]


def compare_workbooks(old_path, new_path):
    """Human-readable report of what changed between two copies of the
    workbook, matched by creditor / council name rather than by row number."""
    old = load_workbook(old_path)
    new = load_workbook(new_path)
    out = [f"OLD: {Path(old_path).name}", f"NEW: {Path(new_path).name}", ""]

    out.append("== Sheets ==")
    for n in old.sheetnames:
        if n not in new.sheetnames:
            out.append(f"  REMOVED sheet: {n}")
    for n in new.sheetnames:
        if n not in old.sheetnames:
            out.append(f"  ADDED sheet: {n}")
    if len(out) == 4:
        out.append("  no sheets added or removed")

    keyed_sheets = [
        (SHEET_GENERAL, GENERAL_COLS, 4),
        (SHEET_COUNCILS, COUNCIL_COLS, 5),
        (SHEET_COUNTY, COUNTY_COLS, 4),
        (SHEET_DIVIDENDS, DIVIDEND_COLS, 3),
    ]
    for sheet, ncols, cmp_cols in keyed_sheets:
        if sheet not in old.sheetnames or sheet not in new.sheetnames:
            continue
        ko, co = _keyed_rows(old[sheet], ncols)
        kn, cn = _keyed_rows(new[sheet], ncols)
        out += ["", f"== {sheet.strip()} ({len(ko)} -> {len(kn)} named rows) =="]
        added = [k for k in kn if k not in ko]
        removed = [k for k in ko if k not in kn]
        changed = []
        for k in kn:
            if k in ko:
                diffs = [(i, ko[k][i], kn[k][i]) for i in range(cmp_cols) if ko[k][i] != kn[k][i]]
                if diffs:
                    changed.append((k, diffs))
        out.append(f"  added {len(added)}, removed {len(removed)}, changed {len(changed)}")
        for k in added:
            out.append(f"  + {' | '.join(_short(c, 100) for c in kn[k][:cmp_cols] if c)}")
        for k in removed:
            out.append(f"  - {' | '.join(_short(c, 100) for c in ko[k][:cmp_cols] if c)}")
        for k, diffs in changed:
            out.append(f"  ~ {kn[k][0]}")
            for i, a, b in diffs:
                out.append(f"      col {i + 1}: {_short(a)!r}")
                out.append(f"           -> {_short(b)!r}")
        for k in kn:
            if k in ko and co.get(k) != cn.get(k):
                out.append(f"  ! {kn[k][0]!r} appears {co.get(k)} time(s) in old, {cn.get(k)} in new")

    if SHEET_WHICH_REP in old.sheetnames and SHEET_WHICH_REP in new.sheetnames:
        out += ["", f"== {SHEET_WHICH_REP} =="]
        wo, wn = old[SHEET_WHICH_REP], new[SHEET_WHICH_REP]
        headings = {}
        for ws in (wo, wn):
            for c in range(1, ws.max_column + 1):
                h = clean_text(ws.cell(1, c).value)
                if h:
                    headings.setdefault(h, {})[ws is wn] = c
        for h, cols in headings.items():
            so = [clean_text(wo.cell(r, cols[False]).value) for r in range(2, wo.max_row + 1)] if False in cols else []
            sn = [clean_text(wn.cell(r, cols[True]).value) for r in range(2, wn.max_row + 1)] if True in cols else []
            so = [x for x in so if x]
            sn = [x for x in sn if x]
            added = [x for x in sn if x not in so]
            removed = [x for x in so if x not in sn]
            out.append(f"  {h}: {len(so)} -> {len(sn)}; added {added or 'none'}; removed {removed or 'none'}")

    for sheet in (SHEET_TIG, SHEET_WATCH, SHEET_TIX, SHEET_EVOLVE):
        if sheet not in old.sheetnames or sheet not in new.sheetnames:
            continue
        to, tn = _text_cells(old[sheet]), _text_cells(new[sheet])
        added = [x for x in tn if x not in to]
        removed = [x for x in to if x not in tn]
        out += ["", f"== {sheet} =="]
        if not added and not removed:
            out.append("  unchanged")
        for x in added:
            out.append(f"  + {_short(x, 300)}")
        for x in removed:
            out.append(f"  - {_short(x, 300)}")

    return "\n".join(out) + "\n"
