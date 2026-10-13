"""
export_excel_criteria / compare_excel_criteria tests.

Builds a tiny workbook with openpyxl, exports it, and checks that
  * the output is in the exact shapes seed_creditor_criteria and
    seed_all_councils parse (round-trip through their parsers),
  * the acknowledged workbook fixes fire and are reported,
  * hand-added rows in the existing markdown are kept and listed,
  * the EVOLVE section is dropped unless asked for,
  * the comparison report matches rows by name.
No database rows are written by any of this.
"""

import datetime
import io
import tempfile
from pathlib import Path

from django.core.management import call_command
from django.test import SimpleTestCase, override_settings
from openpyxl import Workbook

from debt_app.services.excel_criteria import (
    FILE_COUNCILS,
    FILE_COUNTY,
    FILE_DIVIDENDS,
    FILE_GENERAL,
    FILE_WHICH_REP,
    KEPT_MARKER,
    KEPT_SUBHEADING,
    clean_name,
    clean_text,
    compare_workbooks,
    export_workbook,
    split_districts,
)


def _fill(ws, rows):
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if value is not None:
                ws.cell(r, c).value = value


def build_workbook(path, *, with_second_ashford=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "GENERAL CREDITOR"
    _fill(ws, [
        ["Creditor", "Accept / Reject", "Notes / Criteria", "updated criteria", "Phone Number", "Email", "Contact"],
        ["118 118 Money", "Accept", "Accept no mods", None, "02920 893300", "debt@118.com", None],
        [None, None, "a note with no creditor", None, None, None, None],
        ["Startline Motor Finance", "can get a vote if chased", "chase for\nvote", "14/06/24", None, None, None],
        ["Partners Credit Union", "REJECT", "confirmed 12/03/2026 - reject ALL ivas", datetime.datetime(2026, 12, 3), None, None, None],
        ["Weird | Pipe Ltd", "Accept", "cell\xa0with   odd    spacing", datetime.datetime(2024, 9, 9), None, None, None],
        ["Omno  Capital", "DO NOT VOTE", None, None, None, None, None],
        ["Isle\xa0of\xa0Anglesey Finance", "Accept", None, None, None, None, None],
        ["Knowsley Mutual Credit Union", "REJECT", "credit union policy", None, None, None, None],
        ["Hillingdon Credit Union - SEE LONDON COMMUNITY CREDIT UNION", "REJECT", "reject", None, None, None, None],
    ])
    ws = wb.create_sheet("Councils")
    rows = [
        ["Council", "Accept / Reject", "Notes", "Updated Criteria", "Criteria changed from Rej date", "Contact Name", "Contact Number"],
        ["Ashford Borough Council", "REJECT", "reject all IVAs where there are alternative recovery methods", datetime.datetime(2025, 6, 2), None, None, None],
    ]
    if with_second_ashford:
        rows.append(["Ashford Borough Council", "DO NOT VOTE", "last trial 16/09/21 - TIG16827", "8/12/23", None, "Recovery", "01296 585 096"])
    rows += [
        ["Copeland Borough Council FORMS PART OF CUMBERLAND COUNCIL AS OF JAN 2026", "Reject", "always refuse IVAs", None, None, None, None],
        ["Wycombe District Council", "Will Consider", None, None, None, "Council tax", "01494 412 226"],
        ["Wycombe District Council", "Accept", "50P/£ MINIMUM DIV", None, None, None, None],
    ]
    _fill(ws, rows)
    ws = wb.create_sheet("Which Representative")
    _fill(ws, [
        ["TIX", None, "(WATCH) WPM", None, "EVOLVE", None, "EVERYDAY LOANS "],
        ["118 Money", None, "Scottish Power", None, "Mint", None, "GEORGE BANCO"],
        ["Barclays Bank", None, "1:Many", None, "Octopus Energy", None, None],
        [None, None, "Scottish Power", None, None, None, None],   # duplicate in the sheet
        [None, None, "AA (Bank of Ireland) - IVA", None, None, None, None],
    ])
    ws = wb.create_sheet("Dividends ")
    _fill(ws, [
        ["Creditor", "Div Required", "Notes"],
        ["Amigo", None, "DO NOT REQUIRE A DIV NOW"],
        ["Guarantor My Loan", "100P/£", None],
    ])
    ws = wb.create_sheet("County Councils")
    _fill(ws, [
        ["Council", "Accept / Reject", "Notes", "Districts"],
        ["Buckinghamshire County Council", "50P/£ MINIMUM DIV", "Doesnt deal with council tax",
         "South Bucks -                     Chiltern District -  01494 732 077          Wycombe District - 01494 412 226"],
        [None, None, "Dealt with by the district councils", "Sevenoaks District -   01732 227000 - option 5        Borough of Dartford - 01322 343 434"],
    ])
    ws = wb.create_sheet("Banking Groups")
    _fill(ws, [
        ["RBS Banking Group", None, "Barclays"],
        [None, None, None],
        ["Royal Bank of Scotland", None, "Barclays Direct"],
        ["NatWest", None, "Woolwich"],
    ])
    wb.save(str(path))
    return path


EXISTING_GENERAL = """# GENERAL CREDITOR

| Creditor | Accept / Reject | Notes / Criteria | updated criteria | Phone Number | Email | Contact |
| --- | --- | --- | --- | --- | --- | --- |
| 118 118 Money | Accept | old notes |  |  |  |  |
| Anderson Brookes | DO NOT VOTE | Restored 2026-08-18 by hand |  |  |  |  |
| Omno  Capital | DO NOT VOTE | seeded with two spaces |  |  |  |  |
| Knowsley Credit Union | REJECT | old name |  |  |  |  |
| Hillingdon Credit Union - SEE LONDON COMMUNITY CREDIT UNION | REJECT | left by an earlier export |  |  |  |  |
"""

EXISTING_COUNCILS = """# Councils Sheet

| # | Council | Status | Notes | Updated Criteria | Criteria Changed From Rej Date | Contact Name | Contact Number |
|---|---------|--------|-------|------------------|-------------------------------|--------------|----------------|
| 1 | Ashford Borough Council | REJECT | whatever |  |  |  |  |
| 2 | Aylesbury Vale District Council | DO NOT VOTE | last trial 16/09/21 | 8/12/23 |  | Recovery | 01296 585 096 |
| 3 | rth | no voting history |  | 14/2/24 |  |  |  |
"""

EXISTING_WHICH_REP = """# Which Representative

## (WATCH) WPM
- 1:Many
- Monzo

## TIX
- 118 Money
- Lowell

## EVOLVE
- Mint
- Shell Energy

## EVERYDAY LOANS
- GEORGE BANCO
"""


class ExportExcelCriteriaTests(SimpleTestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        self.xlsx = build_workbook(self.root / "criteria.xlsx")
        self.out = self.root / "Excel Criteria"
        self.out.mkdir()
        (self.out / FILE_GENERAL).write_text(EXISTING_GENERAL, encoding="utf-8")
        (self.out / FILE_COUNCILS).write_text(EXISTING_COUNCILS, encoding="utf-8")
        (self.out / FILE_WHICH_REP).write_text(EXISTING_WHICH_REP, encoding="utf-8")

    def tearDown(self):
        self._tmp.cleanup()

    def _export(self, **kw):
        return export_workbook(self.xlsx, self.out, today=datetime.date(2026, 10, 12), **kw)

    # -- cell helpers --------------------------------------------------------

    def test_clean_name_keeps_spacing_and_maps_each_line_break_to_a_space(self):
        self.assertEqual(clean_name("Omno  Capital "), "Omno  Capital")
        self.assertEqual(clean_name("North Yorkshire.\n\n\nNorth Yorkshire Council"), "North Yorkshire.   North Yorkshire Council")
        self.assertEqual(clean_name("Isle\xa0of\xa0Anglesey"), "Isle\xa0of\xa0Anglesey")
        self.assertEqual(clean_name("A | B"), "A / B")

    def test_clean_text_normalises_dates_whitespace_and_pipes(self):
        self.assertEqual(clean_text(datetime.datetime(2026, 12, 3)), "03/12/2026")
        self.assertEqual(clean_text("cell\xa0with   odd \n spacing "), "cell with odd spacing")
        self.assertEqual(clean_text("a | b"), "a / b")
        self.assertEqual(clean_text(None), "")
        self.assertEqual(clean_text(50.0), "50")

    def test_split_districts(self):
        self.assertEqual(
            split_districts("Sevenoaks District -   01732 227000 - option 5        Borough of Dartford - 01322 343 434"),
            ["Sevenoaks District", "01732 227000 - option 5", "Borough of Dartford - 01322 343 434"],
        )

    # -- general creditors ---------------------------------------------------

    def test_general_creditor_fixes_and_kept_rows(self):
        report, outputs = self._export()
        md = outputs[FILE_GENERAL]
        info = report.files[FILE_GENERAL]
        self.assertIn("| Startline Motor Finance | DO NOT VOTE | Sheet says: 'can get a vote if chased'. chase for vote | 14/06/24 |", md)
        self.assertIn("| Partners Credit Union | REJECT | confirmed 12/03/2026 - reject ALL ivas | 12/03/2026 |", md)
        self.assertIn("| Weird / Pipe Ltd | Accept | cell with odd spacing | 09/09/2024 |", md)
        self.assertIn("|  |  | a note with no creditor |", md)  # blank-name rows are preserved (parsers ignore them)
        # name cells keep their exact spacing so they still match the rows seeded from them
        self.assertIn("| Omno  Capital | DO NOT VOTE |  |", md)
        self.assertIn("| Isle\xa0of\xa0Anglesey Finance | Accept |", md)
        self.assertEqual(md.count("Omno  Capital"), 1)  # matched the existing row, not kept twice
        # a remark typed into the name cell goes to the notes; the real name stays
        self.assertIn("| Hillingdon Credit Union | REJECT | See London Community Credit Union. reject |", md)
        # a genuine rename: the new name is written, the old row is not kept, and it is reported
        self.assertIn("| Knowsley Mutual Credit Union | REJECT |", md)
        self.assertNotIn("| Knowsley Credit Union |", md)
        # an uncorrected annotated name left by an earlier export is not kept as a second row
        self.assertEqual(md.count("Hillingdon Credit Union"), 1)
        self.assertEqual(len(info["notes"]), 1)
        self.assertEqual(info["renamed"], ["'Knowsley Credit Union' -> 'Knowsley Mutual Credit Union' (old row not kept; the database sync renames it)"])
        # hand-added row from the existing file survives, after the marker
        self.assertIn(KEPT_MARKER, md)
        self.assertLess(md.index(KEPT_MARKER), md.index("| Anderson Brookes | DO NOT VOTE |"))
        # ...but a row the workbook does have is taken from the workbook, not the old file
        self.assertNotIn("old notes", md)
        self.assertEqual(len(info["fixes"]), 3)
        self.assertEqual(info["kept"], ["'Anderson Brookes' (DO NOT VOTE)"])
        self.assertEqual(info["rows"], 10)

    def test_drop_missing_discards_hand_added_rows(self):
        report, outputs = self._export(drop_missing=True)
        self.assertNotIn("Anderson Brookes", outputs[FILE_GENERAL])
        self.assertNotIn("Aylesbury Vale", outputs[FILE_COUNCILS])
        self.assertNotIn("Lowell", outputs[FILE_WHICH_REP])
        self.assertEqual(report.files[FILE_GENERAL]["kept"], [])

    # -- councils ------------------------------------------------------------

    def test_council_drop_fix_numbering_and_duplicates(self):
        report, outputs = self._export()
        md = outputs[FILE_COUNCILS]
        lines = [l for l in md.splitlines() if l.startswith("| ")]
        self.assertEqual(lines[0], "| # | Council | Status | Notes | Updated Criteria | Criteria Changed From Rej Date | Contact Name | Contact Number |")
        self.assertEqual(md.count("| Ashford Borough Council |"), 1)
        self.assertIn("| 1 | Ashford Borough Council | REJECT |", md)
        self.assertIn("| 2 | Copeland Borough Council | Reject | Forms part of Cumberland Council as of Jan 2026. always refuse IVAs |", md)
        # the duplicate Wycombe 'Will Consider' row is dropped (sheet owner, 2026-10-13); the Accept row stands
        self.assertEqual(md.count("| Wycombe District Council |"), 1)
        self.assertIn("| 3 | Wycombe District Council | Accept | 50P/£ MINIMUM DIV |", md)
        # kept row keeps the numbering going; the junk 'rth' row is not kept
        self.assertIn("| 4 | Aylesbury Vale District Council | DO NOT VOTE |", md)
        self.assertNotIn("| rth |", md)
        self.assertIn("> Total entries: 4", md)
        info = report.files[FILE_COUNCILS]
        self.assertEqual(len(info["dropped"]), 2)
        self.assertIn("Aylesbury Vale", info["dropped"][0])
        self.assertIn("Wycombe", info["dropped"][1])
        self.assertEqual(len(info["fixes"]), 1)

    def test_drop_is_a_no_op_when_workbook_is_already_corrected(self):
        build_workbook(self.xlsx, with_second_ashford=False)
        report, outputs = self._export()
        self.assertEqual([d for d in report.files[FILE_COUNCILS]["dropped"] if "Ashford" in d], [])
        self.assertEqual(outputs[FILE_COUNCILS].count("| Ashford Borough Council |"), 1)

    # -- which representative ------------------------------------------------

    def test_which_representative_order_kept_entries_and_evolve_dropped(self):
        report, outputs = self._export()
        md = outputs[FILE_WHICH_REP]
        headings = [l for l in md.splitlines() if l.startswith("## ")]
        self.assertEqual(headings, ["## (WATCH) WPM", "## TIX", "## EVERYDAY LOANS"])
        self.assertNotIn("EVOLVE", md)
        self.assertNotIn("Octopus Energy", md)
        watch = md[md.index("## (WATCH) WPM"):md.index("## TIX")]
        # sorted case-insensitively and de-duplicated, as the file has always been
        self.assertEqual(
            [l for l in watch.splitlines() if l.startswith("- ")][:3],
            ["- 1:Many", "- AA (Bank of Ireland) - IVA", "- Scottish Power"],
        )
        self.assertEqual(watch.count("- Scottish Power"), 1)
        self.assertIn(KEPT_SUBHEADING, watch)
        self.assertIn("- Monzo", watch)
        tix = md[md.index("## TIX"):md.index("## EVERYDAY LOANS")]
        self.assertIn("- Lowell", tix)
        self.assertEqual(report.files[FILE_WHICH_REP]["kept"], ["(WATCH) WPM: 'Monzo'", "TIX: 'Lowell'"])
        self.assertEqual(len(report.files[FILE_WHICH_REP]["dropped"]), 1)

    def test_keep_evolve_flag(self):
        _, outputs = self._export(keep_evolve=True)
        md = outputs[FILE_WHICH_REP]
        self.assertIn("## EVOLVE", md)
        self.assertIn("- Octopus Energy", md)
        self.assertIn("- Shell Energy", md)  # kept from the existing file

    # -- dividends / county / write ------------------------------------------

    def test_dividends_and_county(self):
        _, outputs = self._export()
        self.assertIn("| Guarantor My Loan | 100P/£ |  |", outputs[FILE_DIVIDENDS])
        self.assertIn("| Amigo |  | DO NOT REQUIRE A DIV NOW |", outputs[FILE_DIVIDENDS])
        county = outputs[FILE_COUNTY]
        self.assertIn("## Buckinghamshire County Council\n\n**Accept / Reject:** 50P/£ MINIMUM DIV  \n**Notes:** Doesnt deal with council tax  \n\n**Districts:**\n- South Bucks\n- Chiltern District - 01494 732 077\n", county)
        self.assertIn("## Kent County Council\n\n**Notes:** Dealt with by the district councils  ", county)

    def test_dry_run_writes_nothing_and_command_writes(self):
        before = (self.out / FILE_GENERAL).read_text(encoding="utf-8")
        out = io.StringIO()
        call_command("export_excel_criteria", "--xlsx", str(self.xlsx), "--out", str(self.out), "--dry-run", stdout=out)
        self.assertEqual((self.out / FILE_GENERAL).read_text(encoding="utf-8"), before)
        self.assertIn("Dry run", out.getvalue())
        self.assertIn("FIX", out.getvalue())

        call_command("export_excel_criteria", "--xlsx", str(self.xlsx), "--out", str(self.out), stdout=io.StringIO())
        self.assertIn("Startline Motor Finance | DO NOT VOTE", (self.out / FILE_GENERAL).read_text(encoding="utf-8"))
        self.assertTrue((self.out / FILE_COUNTY).exists())

    # -- round-trip through the real parsers ---------------------------------

    def test_seeders_parse_the_output(self):
        self._export()
        with override_settings(BASE_DIR=self.root):
            from debt_app.management.commands.seed_creditor_criteria import _parse_strict_sources

            parsed = _parse_strict_sources()
            self.assertEqual(parsed["Startline Motor Finance"]["status"], "DO_NOT_VOTE")
            self.assertEqual(parsed["Partners Credit Union"]["raw_updated_criteria"], "12/03/2026")
            self.assertEqual(parsed["Anderson Brookes"]["status"], "DO_NOT_VOTE")
            self.assertEqual(parsed["Scottish Power"]["representative"], "WATCH")
            self.assertEqual(parsed["Monzo"]["representative"], "WATCH")   # kept entry still maps
            self.assertEqual(parsed["Lowell"]["representative"], "TIX")
            self.assertNotIn("Octopus Energy", parsed)                      # EVOLVE dropped
            self.assertEqual(parsed["Guarantor My Loan"]["min_dividend_pence"], 100)
            self.assertNotIn("a note with no creditor", parsed)

            out = io.StringIO()
            call_command("seed_all_councils", "--dry-run", stdout=out)
            text = out.getvalue()
            self.assertEqual(text.count("Ashford Borough Council -> REJECT"), 1)
            self.assertIn("Copeland Borough Council -> REJECT", text)
            self.assertIn("Aylesbury Vale District Council -> DO_NOT_VOTE", text)
            self.assertNotIn("FORMS PART OF", text)


class CompareWorkbooksTests(SimpleTestCase):
    def test_report_matches_by_name(self):
        with tempfile.TemporaryDirectory() as tmp:
            old = build_workbook(Path(tmp) / "old.xlsx", with_second_ashford=False)
            new = build_workbook(Path(tmp) / "new.xlsx", with_second_ashford=True)
            wb = __import__("openpyxl").load_workbook(str(new))
            ws = wb["GENERAL CREDITOR"]
            ws.insert_rows(2)                       # shift everything down a row
            ws.cell(2, 1).value = "Auto Lend"
            ws.cell(2, 2).value = "Will Consider"
            ws.cell(3, 2).value = "Will Consider"   # 118 118 Money (now row 3): was Accept
            wb["Which Representative"].cell(4, 3).value = "New Watch Co"
            wb.save(str(new))

            text = compare_workbooks(old, new)
            self.assertIn("+ Auto Lend | Will Consider", text)
            self.assertIn("~ 118 118 Money", text)
            self.assertIn("'Accept'", text)
            self.assertIn("'Will Consider'", text)
            self.assertIn("'Ashford Borough Council' appears 1 time(s) in old, 2 in new", text)
            self.assertIn("added ['New Watch Co']", text)
            self.assertNotIn("~ Startline", text)   # shifted rows are not reported as changed

    def test_command_writes_report(self):
        with tempfile.TemporaryDirectory() as tmp:
            old = build_workbook(Path(tmp) / "old.xlsx")
            new = build_workbook(Path(tmp) / "new.xlsx")
            report = Path(tmp) / "report.txt"
            call_command("compare_excel_criteria", "--old", str(old), "--new", str(new), "--out", str(report), stdout=io.StringIO())
            self.assertIn("no sheets added or removed", report.read_text(encoding="utf-8"))
